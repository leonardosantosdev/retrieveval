"""Shared test helpers."""

import pytest


def minimal_pdf(text: str = "") -> bytes:
    """A valid single-page PDF containing ``text``.

    Built by hand rather than with a PDF writer so the tests do not gain a
    dependency they would otherwise never need. Passing an empty string
    produces a page with no extractable text, which is how a scanned document
    behaves as far as this project is concerned.
    """
    content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("ascii") if text else b""
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(content)).encode("ascii") + b" >>\nstream\n"
        + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]

    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += str(number).encode("ascii") + b" 0 obj\n" + body + b"\nendobj\n"

    xref_at = len(out)
    size = str(len(objects) + 1).encode("ascii")
    out += b"xref\n0 " + size + b"\n0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode("ascii")
    out += b"trailer\n<< /Size " + size + b" /Root 1 0 R >>\n"
    out += b"startxref\n" + str(xref_at).encode("ascii") + b"\n%%EOF\n"
    return bytes(out)


@pytest.fixture
def make_pdf():
    return minimal_pdf


def write_docx(path, paragraphs=(), table=None):
    """Write a real .docx at ``path``.

    Unlike the PDF helper this does use ``python-docx`` -- hand-assembling the
    OOXML zip would test the fixture rather than the loader. The package ships
    in the ``dev`` extra for exactly this.

    ``table`` is a list of rows, appended after the paragraphs so that document
    order can be asserted.
    """
    import docx

    document = docx.Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table:
        added = document.add_table(rows=len(table), cols=len(table[0]))
        for row_index, row in enumerate(table):
            for column_index, cell in enumerate(row):
                added.rows[row_index].cells[column_index].text = cell
    document.save(str(path))
    return path


@pytest.fixture
def make_docx():
    return write_docx
